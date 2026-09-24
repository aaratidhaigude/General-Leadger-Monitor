"""
UC14 GL Anomaly Detection Agent — Results Writer
======================================================
"""
from __future__ import annotations
import os, sys, json
from dataclasses import asdict
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1C2B3A"
WHITE = "FFFFFF"
ACTION_FILL = {"AUTO_RECONCILE": "D9F2E3", "DAILY_REPORT": "FDF3D0", "INVESTIGATE": "FBE0E0"}


class ResultsWriter:
    def __init__(self, results_dir: str = None):
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        self.results_dir = results_dir or os.path.join(base_dir, "results")
        try:
            os.makedirs(self.results_dir, exist_ok=True)
        except Exception:
            pass

    def write_results(self, decisions: list) -> dict:
        ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        try:
            base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            results_dir = getattr(self, "results_dir", None) or os.path.join(base_dir, "results")
            os.makedirs(results_dir, exist_ok=True)
            xlsx_path = os.path.join(results_dir, f"uc14_gl_anomaly_run_{ts}.xlsx")
            json_path = os.path.join(results_dir, f"uc14_gl_anomaly_run_{ts}.json")
            self._write_excel(decisions, xlsx_path)
            self._write_json(decisions, json_path, ts)
            return {"excel": xlsx_path, "json": json_path, "run_id": ts, "posting_count": len(decisions)}
        except Exception as e:
            print(f"[ResultsWriter] Notice: Result persistence skipped or restricted ({e}).")
            return {"excel": "", "json": "", "run_id": ts, "posting_count": len(decisions)}


    def _write_excel(self, decisions, path):
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            self._sheet_summary(decisions, writer)
            self._sheet_detailed(decisions, writer)
        self._style_workbook(path)

    def _sheet_summary(self, decisions, writer):
        total = len(decisions)
        by_action = pd.Series([d.action for d in decisions]).value_counts()
        flagged_value = sum(d.amount_inr for d in decisions if d.action != "AUTO_RECONCILE")
        rows = [
            ("Total GL postings processed", total),
            ("Auto-reconciled", int(by_action.get("AUTO_RECONCILE", 0))),
            ("Daily exception report", int(by_action.get("DAILY_REPORT", 0))),
            ("Escalated to investigate", int(by_action.get("INVESTIGATE", 0))),
            ("", ""),
            ("Average ARI", round(sum(d.ari for d in decisions) / total, 1) if total else 0),
            ("Total value flagged (not auto-reconciled), INR", round(flagged_value, 0)),
            ("Run generated (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")),
        ]
        pd.DataFrame(rows, columns=["Metric", "Value"]).to_excel(writer, sheet_name="Summary", index=False)

    def _sheet_detailed(self, decisions, writer):
        df = pd.DataFrame([{
            "Doc No": d.doc_no, "Line": d.line_item, "GL Account": d.gl_account,
            "Doc Type": d.doc_type, "Company": d.company_code, "Posting Date": d.posting_date,
            "Time": d.posting_time, "Amount (INR)": d.amount_inr, "Z-score": d.ari_breakdown.z_score,
            "ARI": d.ari, "Action": d.action, "Justification": d.justification,
            "SAP Action": d.sap_action, "Audit Ref": d.audit_trail_ref,
        } for d in decisions])
        df.to_excel(writer, sheet_name="Detailed Results", index=False)

    def _style_workbook(self, path):
        from openpyxl import load_workbook
        wb = load_workbook(path)
        header_font = Font(bold=True, color=WHITE)
        header_fill = PatternFill("solid", fgColor=NAVY)
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for col_cells in ws.columns:
                length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(45, max(12, length + 2))
            ws.freeze_panes = "A2"
            if ws.title == "Detailed Results":
                action_col = None
                for cell in ws[1]:
                    if cell.value == "Action":
                        action_col = cell.column
                        break
                if action_col:
                    for row in ws.iter_rows(min_row=2, min_col=action_col, max_col=action_col):
                        for cell in row:
                            fill = ACTION_FILL.get(cell.value)
                            if fill:
                                for c in ws[cell.row]:
                                    c.fill = PatternFill("solid", fgColor=fill)
        wb.save(path)

    def _write_json(self, decisions, path, run_id):
        payload = {
            "run_id": run_id, "generated_at_utc": datetime.utcnow().isoformat(),
            "posting_count": len(decisions), "results": [asdict(d) for d in decisions],
        }
        with open(path, "w") as f:
            json.dump(payload, f, indent=2, default=str)
