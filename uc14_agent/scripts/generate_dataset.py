#!/usr/bin/env python3
"""
Generates data/UC14_GLAnomaly_Dataset.xlsx — a 4-section demo dataset for
the UC14 GL Anomaly Detection Agent.

Row placement is driven directly by config/settings.py's SEC_*_HEADER /
SEC_*_START / SEC_*_END constants. Postings are deliberately spread across
all three outcome bands (normal / moderate deviation / severe), the same
lesson carried forward from UC8-UC12's generators.
"""
import os
import sys
import random
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.settings import (
    EXCEL_FILE, SHEET_NAME,
    SEC_A_HEADER, SEC_A_START, SEC_A_END,
    SEC_B_HEADER, SEC_B_START, SEC_B_END,
    SEC_C_HEADER, SEC_C_START, SEC_C_END,
    SEC_D_HEADER, SEC_D_START, SEC_D_END,
)

random.seed(14)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(HERE, "..", EXCEL_FILE)

NAVY = "1C2B3A"
LIGHT = "EEF2F7"
WHITE = "FFFFFF"

COMPANY_CODES = ["IN01", "IN02"]
IC_PARTNERS = ["DE01", "US01", "SG01"]

# (gl_account, gl_desc, doc_type, mean_amount, std_dev_amount, valid_combination)
BASELINES = [
    ("400100", "Domestic Revenue", "DR", 850000, 180000, True),
    ("400200", "Export Revenue", "DR", 1200000, 350000, True),
    ("500200", "Cost of Goods Sold", "SA", 620000, 140000, True),
    ("600100", "Salaries & Wages", "SA", 2100000, 95000, True),
    ("700300", "Travel Expense", "KR", 45000, 22000, True),
    ("700400", "Office Supplies", "KR", 18000, 9000, True),
    ("800100", "Bank Charges", "ZP", 3500, 1800, True),
    ("210500", "Accounts Payable Trade", "KR", 380000, 160000, True),
    ("110200", "Accounts Receivable Trade", "DR", 410000, 175000, True),
    ("999999", "Suspense Account", "SA", 5000, 8000, False),
    ("600200", "Contract Labour", "SA", 310000, 85000, True),
    ("700500", "IT & Software Licenses", "KR", 95000, 40000, True),
    ("800200", "FX Gain/Loss", "SA", 12000, 25000, True),
    ("210600", "Inter-Company Payable", "SA", 500000, 210000, True),
    ("110300", "Inter-Company Receivable", "SA", 500000, 210000, True),
]

DOC_TYPE_LABELS = {"SA": "GL Document", "KR": "Vendor Invoice", "DR": "Customer Invoice", "ZP": "Payment"}


def rdate(base_days_offset, spread=0):
    d = datetime.now() + timedelta(days=base_days_offset + random.randint(0, spread))
    return d.strftime("%Y-%m-%d")


def rtime(off_hours=False):
    if off_hours:
        hour = random.choice([2, 3, 4, 5, 22, 23, 0, 1])
    else:
        hour = random.randint(9, 18)
    minute = random.randint(0, 59)
    return f"{hour:02d}:{minute:02d}"


def section_title(ws, row, span, title):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=title)
    c.font = Font(name="Calibri", bold=True, size=13, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def col_hdr(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor="34486B")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 28


def data_row(ws, row, values, shaded):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name="Calibri", size=10)
        if shaded:
            c.fill = PatternFill("solid", fgColor=LIGHT)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    section_title(ws, 1, 13, "USE CASE 14 — GL ANOMALY DETECTION AGENT")

    # ── Section B: Historical Baseline (written first) ────────────────────
    section_title(ws, SEC_B_HEADER - 1, 5, "SECTION B: HISTORICAL BASELINE PER GL ACCOUNT + DOC TYPE")
    col_hdr(ws, SEC_B_HEADER, ["GL Account", "Doc Type", "Mean Amount", "Std Dev Amount", "Valid Combination"],
            [12, 10, 16, 16, 16])
    n_b = SEC_B_END - SEC_B_START + 1
    for i in range(min(n_b, len(BASELINES))):
        row_no = SEC_B_START + i
        gl, desc, dt, mean, std, valid = BASELINES[i]
        data_row(ws, row_no, [gl, dt, mean, std, "YES" if valid else "NO"], row_no % 2 == 0)

    # ── Section C: Bank Statement Entries ──────────────────────────────────
    section_title(ws, SEC_C_HEADER - 1, 3, "SECTION C: BANK STATEMENT ENTRIES (UTR)")
    col_hdr(ws, SEC_C_HEADER, ["UTR Reference", "Bank Amount", "Bank Date"], [16, 16, 14])
    n_c = SEC_C_END - SEC_C_START + 1
    utrs = []
    for i in range(n_c):
        utr = f"UTR{500000+i}"
        utrs.append(utr)

    # ── Section A: GL Line Items — current period ─────────────────────────
    col_hdr(ws, SEC_A_HEADER, ["Doc No", "Line Item", "GL Account", "Doc Type", "Company Code",
                               "Posting Date", "Posting Time", "Amount (INR)", "Is Intercompany",
                               "IC Partner Company", "IC Partner Amount", "UTR Reference", "Duplicate Flag"],
            [12, 8, 12, 10, 10, 14, 10, 16, 12, 14, 16, 14, 12])

    n_a = SEC_A_END - SEC_A_START + 1
    bank_entries_to_write = []  # (utr, bank_amount, bank_date)
    used_utr_idx = 0

    for i in range(n_a):
        row_no = SEC_A_START + i
        gl, desc, dt, mean, std, valid = random.choice(BASELINES)
        band_roll = random.random()

        duplicate = False
        is_ic = gl in ("210600", "110300")

        if band_roll < 0.50:
            # Normal — should land AUTO_RECONCILE
            amount = round(max(100, random.gauss(mean, std * 0.5)), 0)
            off_hours = False
            severe_extra = None
        elif band_roll < 0.80:
            # Moderate — should land DAILY_REPORT
            amount = round(max(100, mean + random.choice([1, -1]) * std * random.uniform(1.2, 2.3)), 0)
            off_hours = random.random() < 0.4
            severe_extra = None
        else:
            # Severe — should land INVESTIGATE
            severe_kind = random.choice(["big_zscore", "over_1cr", "duplicate", "ic_mismatch"])
            if severe_kind == "over_1cr":
                amount = round(random.uniform(10_500_000, 25_000_000), 0)
            elif severe_kind == "duplicate":
                amount = round(max(100, random.gauss(mean, std * 0.5)), 0)
                duplicate = True
            else:
                amount = round(max(100, mean + random.choice([1, -1]) * std * random.uniform(3.2, 5.0)), 0)
            off_hours = random.random() < 0.5
            severe_extra = severe_kind

        posting_time = rtime(off_hours)
        company = random.choice(COMPANY_CODES)

        ic_partner_company = ""
        ic_partner_amount = ""
        if is_ic:
            ic_partner_company = random.choice(IC_PARTNERS)
            if severe_extra == "ic_mismatch":
                ic_partner_amount = round(amount * random.uniform(1.15, 1.4), 0)
            else:
                ic_partner_amount = amount

        utr_ref = ""
        if not is_ic and random.random() < 0.5 and used_utr_idx < n_c:
            utr_ref = utrs[used_utr_idx]
            # Most bank entries match exactly (clean auto-reconcile case);
            # a couple deliberately don't, to exercise "bank match pending".
            bank_amount = amount if random.random() < 0.85 else round(amount * random.uniform(1.02, 1.08), 0)
            bank_entries_to_write.append((utr_ref, bank_amount, rdate(-random.randint(0, 3))))
            used_utr_idx += 1

        vals = [
            f"{5000000+i}", "1", gl, dt, company, rdate(-random.randint(0, 12)), posting_time,
            amount, "YES" if is_ic else "NO", ic_partner_company, ic_partner_amount, utr_ref,
            "YES" if duplicate else "NO",
        ]
        data_row(ws, row_no, vals, row_no % 2 == 0)

    # Now write Section C bank entries actually referenced by Section A
    for i, (utr, bank_amount, bank_date) in enumerate(bank_entries_to_write[:n_c]):
        row_no = SEC_C_START + i
        data_row(ws, row_no, [utr, bank_amount, bank_date], row_no % 2 == 0)

    # ── Section D: Agent Output Reference (historical outcomes) ───────────
    section_title(ws, SEC_D_HEADER - 1, 4,
                  "SECTION D: AGENT OUTPUT REFERENCE (prior-period actuals, for Learn-phase comparison — not agent input)")
    col_hdr(ws, SEC_D_HEADER, ["Doc No (Historical)", "GL Account", "Final Action Taken", "Resolution Notes"],
            [16, 12, 16, 30])
    n_d = SEC_D_END - SEC_D_START + 1
    for i in range(n_d):
        row_no = SEC_D_START + i
        gl, desc, dt, mean, std, valid = random.choice(BASELINES)
        action = random.choices(["AUTO_RECONCILE", "DAILY_REPORT", "INVESTIGATE"], weights=[0.5, 0.3, 0.2])[0]
        note = {"AUTO_RECONCILE": "Bank match confirmed same day", "DAILY_REPORT": "Reviewed, confirmed valid",
                "INVESTIGATE": "Finance Manager confirmed correction posted"}[action]
        data_row(ws, row_no, [f"{4900000+i}", gl, action, note], row_no % 2 == 0)

    ws.freeze_panes = f"A{SEC_A_HEADER + 1}"
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)

    print(f"Wrote {OUT_PATH}")
    print(f"  Section A (GL Postings):       {n_a} rows (rows {SEC_A_START}-{SEC_A_END})")
    print(f"  Section B (Baselines):         {min(n_b, len(BASELINES))} rows (rows {SEC_B_START}-{SEC_B_END})")
    print(f"  Section C (Bank Entries):      {len(bank_entries_to_write[:n_c])} rows written (rows {SEC_C_START}-{SEC_C_END})")
    print(f"  Section D (Output Reference):  {n_d} rows (rows {SEC_D_START}-{SEC_D_END})")


if __name__ == "__main__":
    main()
